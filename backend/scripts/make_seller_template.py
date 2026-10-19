"""Make the downloadable order sheet for sellers: an Excel file and a CSV with the same columns as
the Add Order form. Run it again after changing the columns; the output goes to frontend/public/.

Usage (from backend/):
    .venv/Scripts/python.exe scripts/make_seller_template.py

The Excel file has drop-down lists for payment type and channel, a date format, and an
"Instructions" sheet with an example. The Orders sheet itself is EMPTY (only the header): an example
row there could be imported by mistake.
"""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

COLUMNS = [
    # (header, width, English help, example)
    ("customer_name", 22, "The customer's name.", "Ayesha Khan"),
    ("phone", 18, "Mobile number: 0300 1234567 or +92 300 1234567. Used to recognise returning customers.", "0300 1234567"),
    ("city", 16, "City the order is going to.", "Lahore"),
    ("products", 34, "Product names in ONE cell, separated by commas.", "Red lawn suit, Dupatta"),
    ("amount", 12, "Order amount in rupees, numbers only.", "4500"),
    ("payment_type", 14, "COD or prepaid (pick from the list).", "cod"),
    ("channel", 14, "Where the order came from (pick from the list).", "instagram"),
    ("order_date", 14, "Date of the order, like 18/10/2026.", "18/10/2026"),
    ("note", 30, "Optional note.", "Call before delivery"),
]
ROWS = 2000                      # rows with drop-downs and date format
OUT = Path(__file__).resolve().parents[2] / "frontend" / "public"


def build_workbook() -> Workbook:
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Orders"
    header_fill = PatternFill("solid", fgColor="0F172A")
    for col, (name, width, _, _) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=col, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
        sheet.column_dimensions[cell.column_letter].width = width
    sheet.freeze_panes = "A2"

    letter = {name: sheet.cell(row=1, column=i).column_letter for i, (name, *_) in enumerate(COLUMNS, start=1)}
    last = ROWS + 1

    payment = DataValidation(type="list", formula1='"cod,prepaid"', allow_blank=True, showErrorMessage=True,
                             errorTitle="Payment type", error="Choose cod or prepaid.")
    channel = DataValidation(type="list", formula1='"instagram,whatsapp,facebook,website,other"', allow_blank=True,
                             showErrorMessage=True, errorTitle="Channel", error="Choose a channel from the list.")
    amount = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True,
                            showErrorMessage=True, errorTitle="Amount", error="Enter the amount as a number above 0.")
    for validation, name in ((payment, "payment_type"), (channel, "channel"), (amount, "amount")):
        sheet.add_data_validation(validation)
        validation.add(f"{letter[name]}2:{letter[name]}{last}")
    for row in range(2, last + 1):
        sheet[f"{letter['order_date']}{row}"].number_format = "DD/MM/YYYY"
        sheet[f"{letter['phone']}{row}"].number_format = "@"      # keep the leading 0 of a phone number

    help_sheet = wb.create_sheet("Instructions")
    help_sheet.column_dimensions["A"].width = 18
    help_sheet.column_dimensions["B"].width = 70
    help_sheet.column_dimensions["C"].width = 26
    help_sheet["A1"] = "How to use this sheet"
    help_sheet["A1"].font = Font(bold=True, size=14)
    help_sheet["A2"] = (
        "Type one order per row on the Orders sheet, then save it as a CSV or keep it as .xlsx and upload it "
        "on the Import Data page of CustomerIQ AI. Upload the same sheet again as it grows: orders already "
        "imported are not added twice."
    )
    help_sheet["A2"].alignment = Alignment(wrap_text=True)
    help_sheet.merge_cells("A2:C2")
    help_sheet.row_dimensions[2].height = 48
    help_sheet.append([])
    help_sheet.append(["Column", "What to write", "Example"])
    for cell in help_sheet[4]:
        cell.font = Font(bold=True)
    for name, _, text, example in COLUMNS:
        help_sheet.append([name, text, example])
    help_sheet.append([])
    help_sheet.append(["Needed", "phone (or a customer_id or email), amount and order_date are needed in every row."])
    help_sheet.append(["Urdu", "ہر قطار میں ایک آرڈر لکھیں۔ فون نمبر، رقم اور تاریخ ضروری ہیں۔"])
    return wb


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    build_workbook().save(OUT / "seller_orders_template.xlsx")
    with (OUT / "seller_orders_template.csv").open("w", newline="", encoding="utf-8") as handle:
        csv.writer(handle).writerow([name for name, *_ in COLUMNS])
    print(f"Wrote {OUT / 'seller_orders_template.xlsx'} and {OUT / 'seller_orders_template.csv'}")


if __name__ == "__main__":
    main()
