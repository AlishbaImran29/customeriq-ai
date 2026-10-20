"""Convert the public UCI "Online Retail II" dataset into weekly units per product, for the Season Forecast.

Source: Chen, D. (2019). Online Retail II [Dataset]. UCI Machine Learning Repository.
https://doi.org/10.24432/C5CG6D  (licence: CC BY 4.0 - credit the source when you share results).
Real transactions of a UK online gift shop, 2009-12-01 to 2011-12-09. It is NOT Pakistani, and it has no
Eid sales, so only its Christmas run-up means anything here.

Rules:
- Only positive-quantity product lines count. Credit notes (invoices starting with "C"), postage, fees,
  samples and manual adjustments are left out (same list as convert_uci_online_retail.py).
- The product is its description (spaces tidied, upper case as in the file).
- Units are added up per product per week (weeks start on Monday).
- To keep the file small only the TOP_PRODUCTS best-selling products are kept.
- Every row carries data_source = "public: Online Retail II".

Only about two Christmas run-ups fit inside this data, so the app will say the forecast is not
proven for it. That is the correct result for so little history.

Usage (from backend/): the input is untrusted, so run in isolated mode and pass paths:
    .venv/Scripts/python.exe -I scripts/convert_uci_product_sales.py INPUT.xlsx [OUTPUT.csv]

Default output: ../data/uci_product_sales_weekly.csv (the data/ folder is gitignored).
"""

import csv
import re
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

NON_PRODUCT_CODES = {
    "POST", "DOT", "M", "C2", "BANK CHARGES", "AMAZONFEE", "ADJUST", "ADJUST2",
    "B", "CRUK", "D", "S", "TEST001", "TEST002",
}
TOP_PRODUCTS = 300
SOURCE = "public: Online Retail II"


def monday(day: date) -> date:
    return day - timedelta(days=day.weekday())


def convert(source: Path) -> tuple[list[dict[str, str]], dict[str, int]]:
    weekly: dict[tuple[str, date], int] = defaultdict(int)
    totals: dict[str, int] = defaultdict(int)
    seen_invoice_sheet: dict[str, int] = {}
    stats = defaultdict(int)
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        for sheet_index, sheet in enumerate(workbook.worksheets):
            for invoice, code, description, quantity, moment, price, _customer, _country in sheet.iter_rows(min_row=2, values_only=True):
                stats["line_rows"] += 1
                invoice = str(invoice or "").strip()
                if not invoice or invoice.upper().startswith("C"):
                    continue
                if seen_invoice_sheet.setdefault(invoice, sheet_index) != sheet_index:
                    continue                                    # the two sheets overlap in December 2010
                if str(code).strip().upper() in NON_PRODUCT_CODES or not description:
                    continue
                try:
                    units, unit_price = int(quantity), float(price)
                except (TypeError, ValueError):
                    continue
                if units <= 0 or unit_price <= 0 or moment is None:
                    continue
                day = moment.date() if isinstance(moment, datetime) else moment
                name = re.sub(r"\s+", " ", str(description)).strip()[:120]
                weekly[(name, monday(day))] += units
                totals[name] += units
    finally:
        workbook.close()
    keep = set(sorted(totals, key=lambda p: -totals[p])[:TOP_PRODUCTS])
    rows = [
        {"product": product, "date": week.isoformat(), "units": str(units), "data_source": SOURCE}
        for (product, week), units in sorted(weekly.items(), key=lambda kv: (kv[0][0], kv[0][1]))
        if product in keep
    ]
    stats["products_in_file"] = len(totals)
    stats["products_kept"] = len(keep)
    stats["rows"] = len(rows)
    return rows, dict(stats)


def main(argv: list[str]) -> int:
    if len(argv) < 2:
        print(__doc__)
        return 2
    source = Path(argv[1])
    target = Path(argv[2]) if len(argv) > 2 else Path(__file__).resolve().parents[2] / "data" / "uci_product_sales_weekly.csv"
    rows, stats = convert(source)
    with target.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=["product", "date", "units", "data_source"])
        writer.writeheader()
        writer.writerows(rows)
    print(f"Wrote {stats['rows']:,} product-week rows ({stats['products_kept']} of {stats['products_in_file']:,} products) to {target}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
