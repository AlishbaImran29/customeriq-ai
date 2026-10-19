"""Convert the public UCI "Online Retail II" dataset into a CustomerIQ import file.

Source: Chen, D. (2019). Online Retail II [Dataset]. UCI Machine Learning Repository.
https://doi.org/10.24432/C5CG6D  (licence: CC BY 4.0 — credit the source when you share
results). Real transactions of a UK online gift retailer, 2009-12-01 to 2011-12-09.

The dataset has one row per product line; CustomerIQ wants one row per order. Rules:
- Rows are grouped by invoice number; the order value is the sum of quantity x price over
  product lines (positive quantity and price; postage, fees, samples and manual
  adjustments are left out).
- Invoices starting with "C" are cancellations / credit notes. They are NOT imported as
  orders and are NOT subtracted from the original orders. Instead each one is linked to the
  order it cancelled: the same customer's most recent earlier invoice that contained the same
  product. That order gets outcome "returned" and the credit note's date as outcome_date.
  Credit notes that match no earlier order are counted and left out. All other orders get
  outcome "delivered" (meaning "not cancelled"; the newest orders may still be cancelled
  later, so a model must ignore orders too recent to have been cancelled yet).
  The label source is "proxy: cancellations": it measures cancellations in a UK shop, not
  cash-on-delivery refusals.
- Rows without a Customer ID (guest checkouts) are skipped.
- The workbook's two sheets overlap in early December 2010; each invoice is counted once.
- The dataset has no email addresses, names, categories, or engagement. Customers get the
  ID "PUB-<number>" and a placeholder email "pub-<number>@example.com".
- Amounts are in pounds (GBP). The app shows a $ symbol: read the numbers as pounds.

Usage (from backend/): the input is untrusted, so run in isolated mode and pass paths:
    .venv/Scripts/python.exe -I scripts/convert_uci_online_retail.py INPUT.xlsx [OUTPUT.csv]

Default output: ../data/public_uci_online_retail_ii.csv (the data/ folder is gitignored).
"""

import csv
import sys
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

# Rows that are not merchandise: postage, bank/marketplace fees, samples, manual changes.
NON_PRODUCT_CODES = {
    "POST", "DOT", "M", "C2", "BANK CHARGES", "AMAZONFEE", "ADJUST", "ADJUST2",
    "B", "CRUK", "D", "S", "TEST001", "TEST002",
}
MAX_ROWS = 50_000  # the import limit (backend/app/services/import_files.py)
CHANNEL = "Online retail"
ID_PREFIX = "PUB-"
LABEL_SOURCE = "proxy: cancellations"  # must match LABEL_PROXY in app/models/order.py


def convert(source: Path) -> tuple[list[dict[str, str]], dict[str, int]]:
    stats: dict[str, int] = defaultdict(int)
    owner_sheet: dict[str, int] = {}                 # invoice -> sheet index that first had it
    invoices: dict[str, dict] = {}
    cancellations: list[dict] = []                   # credit-note product lines, linked afterwards

    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        for sheet_index, sheet in enumerate(workbook.worksheets):
            rows = sheet.iter_rows(min_row=2, values_only=True)  # row 1 is the header
            for invoice, stock_code, _desc, quantity, invoice_date, price, customer_id, country in rows:
                stats["line_rows"] += 1
                invoice = str(invoice).strip() if invoice is not None else ""
                if not invoice:
                    continue
                if owner_sheet.setdefault(invoice, sheet_index) != sheet_index:
                    stats["duplicate_rows_in_overlap"] += 1   # same invoice in the other sheet
                    continue
                if invoice.upper().startswith("C"):
                    stats["cancellation_rows"] += 1
                    _note_cancellation(cancellations, stock_code, quantity, invoice_date, customer_id)
                    continue
                if customer_id is None or str(customer_id).strip() == "":
                    stats["guest_rows"] += 1
                    continue
                if str(stock_code).strip().upper() in NON_PRODUCT_CODES:
                    stats["non_product_rows"] += 1
                    continue
                try:
                    line_value = Decimal(str(price)) * Decimal(str(quantity))
                    positive = Decimal(str(price)) > 0 and Decimal(str(quantity)) > 0
                except Exception:
                    stats["unreadable_rows"] += 1
                    continue
                if not positive or invoice_date is None:
                    stats["non_positive_rows"] += 1
                    continue
                moment = _as_datetime(invoice_date)
                order = invoices.setdefault(
                    invoice,
                    {
                        "customer": str(int(float(customer_id))),
                        "date": moment.date(),
                        "moment": moment,
                        "value": Decimal(0),
                        "country": str(country or "").strip(),
                        "items": 0,                    # units bought
                        "products": set(),             # distinct stock codes
                        "returned_on": None,           # date of the credit note that cancelled it
                    },
                )
                order["value"] += line_value
                order["items"] += int(Decimal(str(quantity)))
                order["products"].add(str(stock_code).strip().upper())
    finally:
        workbook.close()

    _link_cancellations(invoices, cancellations, stats)

    output = []
    for invoice, o in sorted(invoices.items(), key=lambda kv: (kv[1]["date"], kv[0])):
        value = o["value"].quantize(Decimal("0.01"))
        if value <= 0:
            stats["zero_value_orders"] += 1
            continue
        output.append(
            {
                "customer_id": f"{ID_PREFIX}{o['customer']}",
                "email": f"pub-{o['customer']}@example.com",
                "order_id": f"{ID_PREFIX}{invoice}",
                "order_date": o["date"].isoformat(),
                "order_value": str(value),
                "country": o["country"],
                "channel": CHANNEL,
                # Order Risk Check facts
                "order_datetime": o["moment"].strftime("%Y-%m-%d %H:%M:%S"),
                "item_count": max(o["items"], 1),
                "distinct_products": len(o["products"]),
                "outcome": "returned" if o["returned_on"] else "delivered",
                "outcome_date": o["returned_on"].isoformat() if o["returned_on"] else "",
                "label_source": LABEL_SOURCE,
            }
        )
        stats["orders_returned"] += bool(o["returned_on"])
    stats["orders"] = len(output)
    stats["customers"] = len({r["customer_id"] for r in output})
    return output, dict(stats)


def _as_datetime(value) -> datetime:
    """The invoice date with its time of day (Excel gives a datetime; text is read as ISO)."""
    if isinstance(value, datetime):
        return value
    return datetime.fromisoformat(str(value)[:19].replace("T", " "))


def _note_cancellation(cancellations: list[dict], stock_code, quantity, invoice_date, customer_id) -> None:
    """Keep one credit-note product line so it can be linked to the order it cancelled.
    Lines without a customer, a real product or a negative quantity cannot be linked."""
    code = str(stock_code).strip().upper()
    try:
        negative = Decimal(str(quantity)) < 0
        customer = str(int(float(customer_id))) if customer_id not in (None, "") else None
        moment = _as_datetime(invoice_date)
    except Exception:
        return
    if customer is None or code in NON_PRODUCT_CODES or not negative:
        return
    cancellations.append({"customer": customer, "stock": code, "moment": moment})


def _link_cancellations(invoices: dict[str, dict], cancellations: list[dict], stats: dict[str, int]) -> None:
    """Mark the order each credit note cancelled: the same customer's most recent EARLIER
    invoice that contained the same product. Several credit notes may hit one order; the
    earliest date counts, because that is when the return first became known."""
    by_customer: dict[str, list[tuple[datetime, str]]] = defaultdict(list)
    for invoice, o in invoices.items():
        by_customer[o["customer"]].append((o["moment"], invoice))
    for credit in sorted(cancellations, key=lambda c: c["moment"]):
        candidates = [
            (moment, invoice)
            for moment, invoice in by_customer.get(credit["customer"], [])
            if moment <= credit["moment"] and credit["stock"] in invoices[invoice]["products"]
        ]
        if not candidates:
            stats["cancellations_unlinked"] += 1
            continue
        order = invoices[max(candidates)[1]]
        day = credit["moment"].date()
        order["returned_on"] = day if order["returned_on"] is None else min(order["returned_on"], day)
        stats["cancellations_linked"] += 1


def main() -> None:
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    source = Path(sys.argv[1])
    default = Path(__file__).resolve().parents[2] / "data" / "public_uci_online_retail_ii.csv"
    target = Path(sys.argv[2]) if len(sys.argv) > 2 else default

    rows, stats = convert(source)
    if len(rows) > MAX_ROWS:
        raise SystemExit(f"{len(rows):,} orders exceed the {MAX_ROWS:,}-row import limit.")
    target.parent.mkdir(parents=True, exist_ok=True)
    with target.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)

    print(f"Wrote {stats['orders']:,} orders for {stats['customers']:,} customers to {target}")
    print(f"Dates: {rows[0]['order_date']} to {rows[-1]['order_date']}")
    for key in ("line_rows", "duplicate_rows_in_overlap", "cancellation_rows", "cancellations_linked",
                "cancellations_unlinked", "orders_returned", "guest_rows", "non_product_rows",
                "non_positive_rows", "unreadable_rows", "zero_value_orders"):
        print(f"  {key.replace('_', ' ')}: {stats.get(key, 0):,}")


if __name__ == "__main__":
    main()
