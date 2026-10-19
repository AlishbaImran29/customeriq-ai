"""Make the downloadable shipments sheet (Excel and CSV) into frontend/public/.

Usage (from backend/):
    .venv/Scripts/python.exe scripts/make_shipment_template.py

The Excel file has a drop-down for the status, a date format, and an "Instructions" sheet. The
Shipments sheet itself is EMPTY (only the header), so nothing can be imported by mistake.
"""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

COLUMNS = [
    # (header, width, help, example)
    ("order_id", 16, "The order number, if you have it. It links the parcel to the order.", "ORD-1001"),
    ("customer_id", 14, "Or the customer id (any one of order_id, customer_id or phone is enough to link).", "CUS-17"),
    ("phone", 16, "Or the customer's mobile number, like 0300 1234567.", "0300 1234567"),
    ("courier_name", 16, "Required. TCS, Leopards, M&P, Trax... written any way you like.", "TCS"),
    ("city", 16, "The city the parcel goes to.", "Lahore"),
    ("booked_date", 14, "Required. The day you handed it to the courier, like 18/10/2026.", "18/10/2026"),
    ("delivered_or_returned_date", 24, "The day it was delivered or came back. Needed for those two statuses.", "22/10/2026"),
    ("status", 14, "Required. Pick from the list: delivered, returned, in_transit or lost.", "delivered"),
    ("return_reason", 28, "Optional. Why it came back, e.g. Customer refused.", "Customer refused"),
    ("cod_amount", 12, "Optional. The cash on delivery amount.", "4500"),
    ("courier_fee", 12, "Optional. What the courier charged.", "250"),
]
ROWS = 2000
OUT = Path(__file__).resolve().parents[2] / "frontend" / "public"


def build_workbook() -> Workbook:
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Shipments"
    fill = PatternFill("solid", fgColor="0F172A")
    for col, (name, width, _, _) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=col, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        sheet.column_dimensions[cell.column_letter].width = width
    sheet.freeze_panes = "A2"
    letter = {name: sheet.cell(row=1, column=i).column_letter for i, (name, *_) in enumerate(COLUMNS, start=1)}
    last = ROWS + 1

    status = DataValidation(type="list", formula1='"delivered,returned,in_transit,lost"', allow_blank=True,
                            showErrorMessage=True, errorTitle="Status", error="Choose delivered, returned, in_transit or lost.")
    sheet.add_data_validation(status)
    status.add(f"{letter['status']}2:{letter['status']}{last}")
    for row in range(2, last + 1):
        for name in ("booked_date", "delivered_or_returned_date"):
            sheet[f"{letter[name]}{row}"].number_format = "DD/MM/YYYY"
        sheet[f"{letter['phone']}{row}"].number_format = "@"

    help_sheet = wb.create_sheet("Instructions")
    help_sheet.column_dimensions["A"].width = 28
    help_sheet.column_dimensions["B"].width = 72
    help_sheet.column_dimensions["C"].width = 18
    help_sheet["A1"] = "How to use this sheet"
    help_sheet["A1"].font = Font(bold=True, size=14)
    help_sheet["A2"] = (
        "Type one parcel per row on the Shipments sheet and upload the file on the Import Data page of CustomerIQ AI. "
        "Upload it again later with the new results: parcels already imported are updated, not added twice."
    )
    help_sheet["A2"].alignment = Alignment(wrap_text=True)
    help_sheet.merge_cells("A2:C2")
    help_sheet.row_dimensions[2].height = 48
    help_sheet.append([])
    help_sheet.append(["Column", "What to write", "Example"])
    for cell in help_sheet[4]:
        cell.font = Font(bold=True)
    for name, _, text, example in COLUMNS:
        help_sheet.append([name, text, example])
    return wb


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    build_workbook().save(OUT / "shipments_template.xlsx")
    with (OUT / "shipments_template.csv").open("w", newline="", encoding="utf-8") as handle:
        csv.writer(handle).writerow([name for name, *_ in COLUMNS])
    print(f"Wrote {OUT / 'shipments_template.xlsx'} and {OUT / 'shipments_template.csv'}")


if __name__ == "__main__":
    main()
