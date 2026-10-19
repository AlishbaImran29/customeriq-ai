"""The spreadsheet import for a seller's own sheet: no order number, "amount", city, products,
payment type, note and DD/MM/YYYY dates. Every name and number here is made up."""

import io
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.models import Customer, Order
from app.services.import_files import parse_file
from app.services.import_validation import KNOWN_COLUMNS, HEADER_ALIASES, validate_rows
from tests.conftest import requires_database
from tests.test_imports_api import upload

SHEET = "customer_name,phone,city,products,amount,payment_type,channel,order_date,note"
PUBLIC = Path(__file__).resolve().parents[2] / "frontend" / "public"


def validate(*lines: str, header: str = SHEET):
    return validate_rows(parse_file("\n".join([header, *lines]).encode(), "csv"))


# --- reading a seller's sheet -------------------------------------------------------------------------

def test_a_seller_sheet_row_is_read_without_an_order_number():
    result = validate("Ayesha Khan,0300-1234567,Lahore,\"Red lawn suit, Dupatta\",\"4,500\",COD,Instagram,18/10/2026,Call first")
    assert result.errors == []
    row = result.rows[0]
    assert (row.first_name, row.customer_id, row.phone_normalised) == ("Ayesha Khan", "+923001234567", "+923001234567")
    assert (row.city, row.products, row.payment_type, row.note) == ("Lahore", "Red lawn suit, Dupatta", "cod", "Call first")
    assert row.order_value == Decimal("4500.00") and row.order_date == date(2026, 10, 18) and row.sales_channel == "instagram"
    assert row.order_id.startswith("SH-") and len(row.order_id) == 15 and row.email is None


def test_amount_is_another_name_for_order_value_and_both_together_are_refused():
    assert validate("0300 1234567,100,2026-01-02", header="phone,amount,order_date").rows[0].order_value == Decimal("100.00")
    from app.core.errors import ApiError
    with pytest.raises(ApiError) as exc:
        validate("0300 1234567,100,100,2026-01-02", header="phone,amount,order_value,order_date")
    assert exc.value.details["duplicateColumns"] == ["order_value"]
    assert HEADER_ALIASES == {"amount": "order_value"}


@pytest.mark.parametrize(
    ("typed", "expected"),
    [("18/10/2026", date(2026, 10, 18)), ("3/4/2026", date(2026, 4, 3)), ("03-04-2026", date(2026, 4, 3)),
     ("03.04.2026", date(2026, 4, 3)), ("2026-04-03", date(2026, 4, 3))],
)
def test_dates_are_day_first_like_in_pakistan_and_iso_still_works(typed, expected):
    assert validate(f"0300 1234567,5,{typed}", header="phone,amount,order_date").rows[0].order_date == expected


@pytest.mark.parametrize("typed", ["31/02/2026", "2026/04/03", "April 3", "13/13/2026", "01/01/2999"])
def test_dates_that_cannot_be_read_are_refused_without_echoing_them(typed):
    result = validate(f"0300 1234567,5,{typed}", header="phone,amount,order_date")
    assert result.rows == [] and [e.column for e in result.errors] == ["order_date"]
    assert typed not in result.errors[0].message


@pytest.mark.parametrize(
    ("typed", "expected"),
    [("COD", "cod"), ("cash on delivery", "cod"), ("Cash", "cod"), ("prepaid", "prepaid"), ("Pre-paid", "prepaid"),
     ("advance", "prepaid"), ("EasyPaisa", "prepaid"), ("jazz cash", "prepaid"), ("bank transfer", "prepaid"), ("", None)],
)
def test_payment_type_from_the_words_sellers_use(typed, expected):
    row = validate(f"0300 1234567,5,2026-01-02,{typed}", header="phone,amount,order_date,payment_type").rows[0]
    assert row.payment_type == expected


@pytest.mark.parametrize("typed", ["unpaid", "maybe", "cod or prepaid"])
def test_an_unclear_payment_type_is_an_error_not_a_guess(typed):
    result = validate(f"0300 1234567,5,2026-01-02,{typed}", header="phone,amount,order_date,payment_type")
    assert result.rows == [] and result.errors[0].column == "payment_type" and typed not in result.errors[0].message


def test_an_explicit_name_column_wins_over_customer_name_and_long_values_are_refused():
    row = validate("0300 1234567,5,2026-01-02,Sana,Sana Khan", header="phone,amount,order_date,first_name,customer_name").rows[0]
    assert row.first_name == "Sana"
    result = validate(f"0300 1234567,5,2026-01-02,{'x' * 1001}", header="phone,amount,order_date,products")
    assert result.errors[0].column == "products"


def test_generated_ids_are_stable_and_identical_rows_get_a_number():
    line = "0300 1234567,4500,2026-01-02,Kurta"
    first = validate(line, line, "0300 1234567,4500,2026-01-02,Scarf", header="phone,amount,order_date,products").rows
    again = validate(line, line, "0300 1234567,4500,2026-01-02,Scarf", header="phone,amount,order_date,products").rows
    assert [r.order_id for r in first] == [r.order_id for r in again]                       # the same sheet gives the same ids
    assert first[1].order_id == first[0].order_id + "-2" and len({r.order_id for r in first}) == 3


def test_explicit_order_ids_are_kept_and_duplicates_still_refused():
    result = validate("O1,0300 1234567,5,2026-01-02", "O1,0300 1234567,6,2026-01-02", header="order_id,phone,amount,order_date")
    assert [r.order_id for r in result.rows] == ["O1"] and result.errors[0].column == "order_id"


# --- the template files -------------------------------------------------------------------------------

def test_the_template_columns_are_all_understood_by_the_importer():
    csv_path, xlsx_path = PUBLIC / "seller_orders_template.csv", PUBLIC / "seller_orders_template.xlsx"
    if not csv_path.exists():
        pytest.skip("frontend folder not available")
    headers = csv_path.read_text(encoding="utf-8").strip().split(",")
    assert headers == SHEET.split(",")
    assert all(HEADER_ALIASES.get(h, h) in KNOWN_COLUMNS for h in headers)
    sheet = load_workbook(xlsx_path)["Orders"]
    assert [c.value for c in sheet[1]] == headers
    assert sheet.max_row >= 2 and all(c.value is None for c in sheet[2])                    # no example row to import by mistake


def test_a_filled_in_excel_template_imports():
    path = PUBLIC / "seller_orders_template.xlsx"
    if not path.exists():
        pytest.skip("frontend folder not available")
    wb = load_workbook(path)
    sheet = wb["Orders"]
    sheet.append([])  # nothing: a blank row must not matter
    for column, value in enumerate(["Ayesha Khan", "0300 1234567", "Lahore", "Lawn suit, Dupatta", 4500, "cod", "whatsapp", date(2026, 10, 18), "x"], start=1):
        sheet.cell(row=2, column=column, value=value)
    buffer = io.BytesIO()
    wb.save(buffer)
    result = validate_rows(parse_file(buffer.getvalue(), "xlsx"))
    assert result.errors == [] and len(result.rows) == 1
    row = result.rows[0]
    assert (row.phone_normalised, row.city, row.payment_type, row.sales_channel, row.order_date) == ("+923001234567", "Lahore", "cod", "whatsapp", date(2026, 10, 18))


# --- saving ---------------------------------------------------------------------------------------------

@requires_database
def test_a_seller_sheet_is_saved_with_the_new_fields(db_client, db_session):
    response = upload(db_client, "\n".join([
        SHEET,
        'Ayesha Khan,0300-1234567,Lahore,"Red lawn suit, Dupatta",4500,COD,Instagram,18/10/2026,Call first',
        "Bilal,0321 9876543,Karachi,Kurta,2500,prepaid,WhatsApp,19/10/2026,",
    ]))
    assert response.status_code == 201, response.text
    assert response.json()["rowsImported"] == 2 and response.json()["rowsRejected"] == 0
    orders = {o.customer.first_name: o for o in db_session.scalars(select(Order))}
    ayesha = orders["Ayesha Khan"]
    assert ayesha.external_id.startswith("SH-") and (ayesha.city, ayesha.products, ayesha.payment_type, ayesha.note) == ("Lahore", "Red lawn suit, Dupatta", "cod", "Call first")
    assert (ayesha.sales_channel, ayesha.customer.phone_normalised, ayesha.customer.email) == ("instagram", "+923001234567", None)
    assert orders["Bilal"].payment_type == "prepaid" and orders["Bilal"].order_date == date(2026, 10, 19)


@requires_database
def test_uploading_a_growing_sheet_again_adds_only_the_new_rows(db_client, db_session):
    first = ["Ayesha,0300 1234567,Lahore,Suit,4500,cod,instagram,18/10/2026,", "Bilal,0321 9876543,Karachi,Kurta,2500,cod,whatsapp,19/10/2026,"]
    assert upload(db_client, "\n".join([SHEET, *first])).json()["rowsImported"] == 2
    again = upload(db_client, "\n".join([SHEET, *first, "Zara,0345 1112223,Multan,Scarf,900,cod,facebook,10/10/2026,"])).json()
    assert (again["rowsImported"], again["rowsUpdated"]) == (1, 2)
    assert db_session.scalar(select(func.count()).select_from(Order)) == 3
    assert db_session.scalar(select(func.count()).select_from(Customer)) == 3


@requires_database
def test_a_blank_cell_in_a_later_upload_does_not_erase_what_is_saved(db_client, db_session):
    upload(db_client, "\n".join([SHEET, "Ayesha,0300 1234567,Lahore,Suit,4500,cod,instagram,18/10/2026,Call first"]))
    upload(db_client, "\n".join([SHEET, "Ayesha,0300 1234567,Lahore,Suit,4500,,,18/10/2026,"]))
    order = db_session.scalar(select(Order))
    assert (order.payment_type, order.sales_channel, order.note) == ("cod", "instagram", "Call first")
