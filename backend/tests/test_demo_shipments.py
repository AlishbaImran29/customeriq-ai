"""The made-up demo shipments, the shipment template files, and the minimum-count rule.

The orders used here are invented too: this file never touches real data."""

import io
import random
from collections import Counter, defaultdict
from datetime import date, timedelta
from importlib import util
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.models.shipment import SOURCE_DEMO
from app.services import data_rules
from app.services.import_files import parse_file
from app.services.names import normalise_city, normalise_courier
from app.services.shipment_import import KNOWN_COLUMNS, validate_shipments

ROOT = Path(__file__).resolve().parents[2]
spec = util.spec_from_file_location("generate_demo_shipments", ROOT / "backend" / "scripts" / "generate_demo_shipments.py")
demo = util.module_from_spec(spec)
spec.loader.exec_module(demo)


def made_up_orders(n: int = 9000):
    rng = random.Random(1)
    return [(f"X-{i}", f"C-{rng.randrange(2000)}", date(2026, 1, 1) + timedelta(days=rng.randrange(0, 240)), round(rng.uniform(500, 9000), 2)) for i in range(n)]


@pytest.fixture(scope="module")
def rows():
    return demo.generate(made_up_orders())


def test_the_generator_always_marks_its_rows_as_synthetic(rows):
    assert demo.SOURCE == SOURCE_DEMO == "demo (synthetic)"
    assert {r["data_source"] for r in rows} == {SOURCE_DEMO} and len(rows) == 9000


def test_the_same_input_gives_the_same_rows(rows):
    assert demo.generate(made_up_orders()) == rows
    assert demo.generate(made_up_orders(), seed=43) != rows


def test_every_generated_row_passes_the_real_shipment_validation(rows):
    import csv

    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=demo.COLUMNS)
    writer.writeheader()
    writer.writerows(rows)
    result = validate_shipments(parse_file(buffer.getvalue().encode(), "csv"))
    assert result.errors == [] and len(result.rows) == len(rows)


def test_messy_spellings_are_in_the_file_and_all_become_the_standard_names(rows):
    couriers = {r["courier_name"] for r in rows}
    assert {"t.c.s", "post ex", "LEOPARDS"} <= couriers                               # messy on purpose
    assert {normalise_courier(c).name for c in couriers} == set(demo.COURIERS)
    cities = {r["city"] for r in rows}
    assert {"Isb", "pindi", "khi"} <= cities
    assert {normalise_city(c).name for c in cities} == set(demo.CITIES)


def test_the_planted_pattern_is_there_for_the_unusual_returns_alert_to_find(rows):
    def rate(pick):
        closed = [r for r in rows if r["status"] in ("delivered", "returned") and pick(r)]
        return sum(r["status"] == "returned" for r in closed) / len(closed), len(closed)

    in_multan = lambda r: normalise_city(r["city"]).name == "Multan"                   # noqa: E731
    leopards = lambda r: normalise_courier(r["courier_name"]).name == "Leopards"        # noqa: E731
    planted, planted_n = rate(lambda r: in_multan(r) and leopards(r))
    others, others_n = rate(lambda r: in_multan(r) and not leopards(r))
    assert planted > 1.6 * others and planted_n >= data_rules.MIN_GROUP_COUNT and others_n >= data_rules.MIN_GROUP_COUNT


def test_the_last_days_are_still_in_transit_and_dates_make_sense(rows):
    last = max(date.fromisoformat(r["booked_date"]) for r in rows)
    recent = [r for r in rows if date.fromisoformat(r["booked_date"]) > last - timedelta(days=demo.IN_TRANSIT_DAYS)]
    assert recent and all(r["status"] == "in_transit" for r in recent)
    for r in rows:
        if r["status"] in ("delivered", "returned"):
            assert date.fromisoformat(r["delivered_or_returned_date"]) >= date.fromisoformat(r["booked_date"])
        else:
            assert r["delivered_or_returned_date"] == ""
    assert Counter(r["status"] for r in rows)["lost"] > 0


def test_return_rates_differ_by_city_the_way_the_recipe_says(rows):
    by_city = defaultdict(lambda: [0, 0])
    for r in rows:
        if r["status"] in ("delivered", "returned"):
            c = normalise_city(r["city"]).name
            by_city[c][0] += 1
            by_city[c][1] += r["status"] == "returned"
    assert by_city["Quetta"][1] / by_city["Quetta"][0] > by_city["Islamabad"][1] / by_city["Islamabad"][0]


# --- the template files ------------------------------------------------------------------------------------

def test_the_shipment_templates_match_what_the_importer_reads():
    csv_path, xlsx_path = ROOT / "frontend" / "public" / "shipments_template.csv", ROOT / "frontend" / "public" / "shipments_template.xlsx"
    if not csv_path.exists():
        pytest.skip("frontend folder not available")
    headers = csv_path.read_text(encoding="utf-8").strip().split(",")
    assert all(h in KNOWN_COLUMNS for h in headers) and {"courier_name", "booked_date", "status"} <= set(headers)
    sheet = load_workbook(xlsx_path)["Shipments"]
    assert [c.value for c in sheet[1]] == headers and all(c.value is None for c in sheet[2])     # no example row to import by mistake


def test_a_filled_in_excel_template_imports():
    path = ROOT / "frontend" / "public" / "shipments_template.xlsx"
    if not path.exists():
        pytest.skip("frontend folder not available")
    wb = load_workbook(path)
    sheet = wb["Shipments"]
    values = ["O-1", None, None, "t.c.s", "isb", date(2026, 9, 1), date(2026, 9, 4), "delivered", None, 4500, 250]
    for column, value in enumerate(values, start=1):
        sheet.cell(row=2, column=column, value=value)
    buffer = io.BytesIO()
    wb.save(buffer)
    result = validate_shipments(parse_file(buffer.getvalue(), "xlsx"))
    assert result.errors == [] and (result.rows[0].courier.name, result.rows[0].city.name) == ("TCS", "Islamabad")


# --- the minimum-count rule -------------------------------------------------------------------------------

def test_the_minimum_count_lives_in_one_place_and_defaults_to_20():
    assert data_rules.MIN_GROUP_COUNT == 20
    assert data_rules.enough(20) and not data_rules.enough(19)
    assert data_rules.more_needed(5) == 15 and data_rules.more_needed(20) == 0 and data_rules.more_needed(500) == 0


def test_changing_the_number_changes_the_rule(monkeypatch):
    monkeypatch.setattr(data_rules, "MIN_GROUP_COUNT", 50)
    assert not data_rules.enough(30) and data_rules.more_needed(30) == 20
